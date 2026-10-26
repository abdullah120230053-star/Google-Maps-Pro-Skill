"""Offline behavior tests. Run: python -m unittest discover -s scripts -p 'test_*.py'."""
import csv
import io
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import unittest
import urllib.error
import zipfile
from contextlib import redirect_stdout, redirect_stderr
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from unittest import mock
from xml.etree import ElementTree as ET

import campaign
import pro
import scrape


class FiltersAndScores(unittest.TestCase):
    def setUp(self):
        self.row = {"title": "عيادة النور", "phone": "+20 100", "emails": "hello@clinic.test",
                    "website": "https://clinic.test", "category": "Dental Clinic", "status": "OPERATIONAL",
                    "review_rating": "٤٫٥", "review_count": "١٬٢٠٠", "instagram": "https://instagram.com/clinic"}

    def test_filters_boundaries_unknown_and_arabic_numbers(self):
        self.assertIsNone(pro.filter_reason(self.row, {"min_rating": 4.5, "max_rating": 4.5, "min_reviews": 1200, "max_reviews": 1200}))
        self.assertEqual(pro.filter_reason(dict(self.row, review_rating=""), {"min_rating": 0}), "min_rating")
        self.assertEqual(pro.filter_reason(self.row, {"max_reviews": 100}), "max_reviews")
        self.assertIsNone(pro.number("nan"))

    def test_categories_contacts_and_status(self):
        self.assertIsNone(pro.filter_reason(self.row, {"category": ["dentAL"], "has_email": True,
                                                     "has_phone": True, "has_website": True, "has_socials": True,
                                                     "business_status": ["open"]}))
        for settings, reason in (({"exclude_category": ["clinic"]}, "exclude_category"),
                                 ({"no_website": True}, "no_website"),
                                 ({"business_status": ["permanently closed"]}, "business_status")):
            self.assertEqual(pro.filter_reason(self.row, settings), reason)
        self.assertIsNone(pro.filter_reason({"status": "مغلق نهائيًا"}, {"business_status": ["permanently closed"]}))

    def test_balanced_score_is_explainable_and_capped(self):
        pro.score_lead(self.row, pro.scoring_rules())
        self.assertEqual(self.row["lead_score"], 100)
        self.assertIn("+20 phone available", self.row["score_reasons"])
        self.assertEqual(pro.filter_reason(self.row, {"min_score": 100}), None)
        pro.score_lead({}, pro.scoring_rules())
        rules = pro.scoring_rules(overrides={"weights": {"phone": 100}})
        self.assertEqual(pro.score_lead(self.row, rules)["lead_score"], 100)

    def test_website_profile_distinguishes_unknown_from_failure(self):
        rules = pro.scoring_rules("website-services")
        blocked = pro.score_lead({"website": "https://clinic.test", "website_status": "blocked", "website_https": ""}, rules)
        self.assertEqual(blocked["lead_score"], 0)
        unreachable = pro.score_lead({"website": "https://clinic.test", "website_status": "unreachable"}, rules)
        self.assertEqual(unreachable["lead_score"], 25)
        self.assertEqual(pro.score_lead({}, rules)["lead_score"], 35)

    def test_invalid_scoring_rejected(self):
        for overrides in ({"weights": {"typo": 10}}, {"weights": {"email": -1}}, {"weights": {"email": True}}, {"rating_min": 6}):
            with self.assertRaises(ValueError):
                pro.scoring_rules(overrides=overrides)


class Discovery(unittest.TestCase):
    def test_poll_failure_preserves_job_id_and_active_state(self):
        with mock.patch.object(scrape, "req", side_effect=[(201, b'{"id":"active-123"}'), urllib.error.URLError("connection lost")]), redirect_stdout(io.StringIO()):
            with self.assertRaises(scrape.JobError) as result:
                scrape.run_job({"max_time": 600})
        self.assertEqual(result.exception.job_id, "active-123")
        self.assertTrue(result.exception.active_unknown)

    def test_real_local_http_redirect_unicode_and_shared_cache(self):
        counts = {}
        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_GET(self):
                counts[self.path] = counts.get(self.path, 0) + 1
                if self.path == "/redirect":
                    self.send_response(302)
                    self.send_header("Location", "/")
                    self.end_headers()
                    return
                code = 403 if self.path == "/blocked" else 404 if self.path == "/missing" else 200
                data = '<a href="/اتصل">تواصل</a><a href="https://instagram.com/clinic">IG</a> hello@clinic.test'.encode("utf-8")
                self.send_response(code)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)
        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        base = f"http://127.0.0.1:{server.server_port}"
        scrape._fetch_page.cache_clear()
        try:
            rows = [{"website": base + "/"}]
            with redirect_stdout(io.StringIO()):
                scrape.enrich_emails(rows)
                scrape.enrich_socials(rows)
                scrape.check_websites(rows)
            self.assertEqual(counts["/"], 1)
            self.assertEqual(rows[0]["emails"], "hello@clinic.test")
            self.assertEqual(rows[0]["instagram_source"], base + "/")
            self.assertIn("%D8", scrape._fetch_page(base + "/اتصل")["url"])
            self.assertEqual(scrape._fetch_page(base + "/redirect")["url"], base + "/")
            self.assertEqual(scrape._fetch_page(base + "/blocked")["state"], "blocked")
            self.assertEqual(scrape._fetch_page(base + "/missing")["state"], "http_error")
            self.assertEqual(scrape._fetch_page("http://[invalid")["state"], "invalid_url")
        finally:
            server.shutdown()
            server.server_close()
            scrape._fetch_page.cache_clear()

    def test_direct_social_profile_is_not_crawled(self):
        rows = [{"website": "https://facebook.com/clinic"}]
        scrape._fetch_page.cache_clear()
        with mock.patch.object(scrape.urllib.request, "urlopen") as request, redirect_stdout(io.StringIO()):
            scrape.enrich_emails(rows)
            scrape.enrich_socials(rows)
            scrape.check_websites(rows)
            request.assert_not_called()
        self.assertEqual(rows[0]["facebook"], "https://facebook.com/clinic")
        self.assertEqual(rows[0]["facebook_source"], "https://facebook.com/clinic")
        self.assertEqual(rows[0]["website_status"], "social_profile")
        scrape._fetch_page.cache_clear()

    def test_socials_profiles_not_posts_and_no_fake_domains(self):
        text = '''<a href="https://instagram.com/p/post">post</a>
        https://instagram.com/clinic/?utm_source=site
        https://facebook.com/sharer/sharer.php
        https://facebook.com/profile.php?id=123456
        https://eg.linkedin.com/company/clinic/posts/
        https://www.youtube.com/@clinic/videos
        https://www.tiktok.com/@clinic/video/123
        https://twitter.com/clinic/status/123
        https://api.whatsapp.com/send?phone=%2B201001234567&amp;text=Hi
        https://instagram.com.evil.test/fake'''
        links = scrape._find_socials(text)
        self.assertEqual(links, {"instagram": "https://instagram.com/clinic", "facebook": "https://facebook.com/profile.php?id=123456",
                                 "linkedin": "https://linkedin.com/company/clinic", "youtube": "https://youtube.com/@clinic",
                                 "tiktok": "https://tiktok.com/@clinic", "x": "https://x.com/clinic", "whatsapp": "https://wa.me/201001234567"})
        self.assertFalse(any(scrape._find_socials("https://instagram.com.evil.test/fake").values()))

    def test_jsonld_escaped_and_protocol_relative_links(self):
        links = scrape._find_socials(r'{"sameAs":["https:\/\/instagram.com\/clinic", "//facebook.com/clinic"]}')
        self.assertEqual(links["instagram"], "https://instagram.com/clinic")
        self.assertEqual(links["facebook"], "https://facebook.com/clinic")

    def test_arabic_contact_links_and_entity_emails(self):
        links = scrape.contact_links('<a href="/ar/اتصل-بنا">اتصل بنا</a><a href="/team">من نحن</a>', "https://clinic.test")
        self.assertIn("https://clinic.test/ar/اتصل-بنا", links)
        self.assertIn("https://clinic.test/team", links)
        self.assertEqual(scrape.extract_emails("hello&#64;clinic.test", "clinic.test"), ["hello@clinic.test"])

    def test_social_source_and_cached_homepage(self):
        rows = [{"website": "https://clinic.test"}]
        page = {"html": '<a href="https://instagram.com/clinic">IG</a>', "url": "https://www.clinic.test/", "status": 200, "state": "reachable"}
        with mock.patch.object(scrape, "_fetch_page", return_value=page), redirect_stdout(io.StringIO()):
            scrape.enrich_socials(rows)
        self.assertEqual(rows[0]["instagram_source"], "https://www.clinic.test/")

    def test_website_observations(self):
        pages = {
            "": {"html": "", "url": "", "status": None, "state": "missing"},
            "https://blocked.test": {"html": "", "url": "https://blocked.test", "status": 403, "state": "blocked"},
            "http://clinic.test": {"html": '<a href="/اتصل">تواصل معنا</a>', "url": "http://clinic.test", "status": 200, "state": "reachable"}}
        rows = [{"website": key} for key in pages]
        with mock.patch.object(scrape, "_fetch_page", side_effect=lambda url: pages[url]), redirect_stdout(io.StringIO()):
            scrape.check_websites(rows)
        self.assertEqual(rows[0]["website_status"], "missing")
        self.assertEqual(rows[1]["website_https"], "")
        self.assertIn("manual review", rows[1]["website_opportunities"])
        self.assertFalse(rows[2]["website_https"])
        self.assertEqual(rows[2]["website_contact_link"], "http://clinic.test/اتصل")
        self.assertNotIn("no contact", rows[2]["website_opportunities"])


class Campaigns(unittest.TestCase):
    def write(self, root, name, value):
        path = Path(root, name)
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        return str(path)

    def test_campaign_overrides_and_localization(self):
        with tempfile.TemporaryDirectory() as root:
            cfg = self.write(root, "campaign.json", {"keywords": ["dentists in {city}"],
                "cities": [{"name": "Cairo", "names": {"ar": "القاهرة"}, "lat": 30, "lon": 31}],
                "languages": ["ar"], "filters": {"category": ["clinic"], "no_website": True},
                "options": {"depth": 4, "socials": False}})
            a, _, targets, _ = campaign.parse_args(["--campaign", cfg, "--lang", "en", "--category", "dentist", "--has-website", "--socials", "--depth", "2"])
            self.assertEqual(a.lang, ["en"])
            self.assertEqual(a.category, ["dentist"])
            self.assertFalse(a.no_website)
            self.assertTrue(a.socials)
            self.assertEqual(a.depth, 2)
            self.assertEqual(campaign.localized_keywords(["عيادات في {city}"], targets[0], "ar", False), ["عيادات في القاهرة"])

    def test_bad_campaign_and_args_fail_before_requests(self):
        bad = [(["x", "--min-rating", "nan"]), (["x", "--min-rating", "5", "--max-rating", "2"]),
               (["x", "--no-socials", "--has-socials"]), (["x", "--out", "x.xlsx", "--json"]),
               (["dentists in Cairo", "--city", "Cairo", "--city", "Giza"]), (["x", "--fields", ", "])]
        with mock.patch.object(scrape, "req") as request, redirect_stderr(io.StringIO()):
            for argv in bad:
                with self.subTest(argv=argv), self.assertRaises(SystemExit) as result:
                    scrape.main(argv)
                self.assertEqual(result.exception.code, 2)
            request.assert_not_called()
        with tempfile.TemporaryDirectory() as root:
            for value in ({"options": {"socails": True}}, {"options": {"socials": "false"}}, {"cities": [{"name": "Cairo", "lat": 30}]}):
                path = self.write(root, "bad.json", value)
                with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                    scrape.main(["--campaign", path, "dentists"])

    def test_multicity_bilingual_filters_scores_and_report(self):
        with tempfile.TemporaryDirectory() as root:
            path = self.write(root, "campaign.json", {"name": "Egypt dentists", "cities": ["Cairo", "Giza"],
                "keywords": ["dentists in {city}"], "keywords_by_language": {"ar": ["عيادات أسنان في {city}"]},
                "languages": ["en", "ar"], "filters": {"min_reviews": 30, "has_phone": True},
                "options": {"deep_emails": False, "socials": False}})
            def job(body):
                return "job-" + body["lang"], [
                    {"place_id": "P1", "title": "عيادة النور", "phone": "+20 100", "review_count": "٤٠", "review_rating": "4.5"},
                    {"place_id": "P2", "title": "New clinic", "phone": "+20 101", "review_count": "2"}]
            out = str(Path(root, "leads.json"))
            with mock.patch.object(scrape, "req"), mock.patch.object(scrape, "geocode", return_value=("30", "31")), \
                 mock.patch.object(scrape, "run_job", side_effect=job) as jobs, redirect_stdout(io.StringIO()):
                report = scrape.main(["--campaign", path, "--out", out])
            self.assertEqual(jobs.call_count, 4)
            payloads = [call.args[0] for call in jobs.call_args_list]
            self.assertIn("عيادات أسنان في Giza", [b["keywords"][0] for b in payloads])
            self.assertEqual(report["totals"]["raw_rows"], 8)
            self.assertEqual(report["totals"]["unique_collected"], 2)
            self.assertEqual(report["totals"]["duplicates_removed"], 6)
            self.assertEqual(report["totals"]["filtered_out"], 1)
            data = json.loads(Path(out).read_text(encoding="utf-8"))
            self.assertEqual(data[0]["title"], "عيادة النور")
            self.assertEqual(data[0]["campaign_cities"], "Cairo; Giza")
            self.assertEqual(data[0]["search_languages"], "en; ar")
            self.assertEqual(data[0]["lead_score"], 50)
            self.assertEqual(report["locations"][1]["exported_rows"], 1)
            self.assertTrue(Path(root, "leads.report.json").exists())

    def test_custom_projection_does_not_break_enrichment_or_filters(self):
        with tempfile.TemporaryDirectory() as root:
            out = str(Path(root, "leads.csv"))
            def enrich(rows):
                self.assertIn("website", rows[0])
                rows[0]["emails"] = "hello@clinic.test"
            def socials(rows):
                self.assertIn("website", rows[0])
                rows[0]["instagram"] = "https://instagram.com/clinic"
            with mock.patch.object(scrape, "req"), mock.patch.object(scrape, "run_job", return_value=("job", [{"title": "عيادة", "website": "https://clinic.test"}])), \
                 mock.patch.object(scrape, "enrich_emails", side_effect=enrich), mock.patch.object(scrape, "enrich_socials", side_effect=socials), redirect_stdout(io.StringIO()):
                report = scrape.main(["dentists", "30", "31", "--fields", "title,emails,instagram", "--has-email", "--has-socials", "--out", out])
            with open(out, encoding="utf-8-sig", newline="") as target:
                rows = list(csv.DictReader(target))
            self.assertEqual(rows[0]["instagram"], "https://instagram.com/clinic")
            self.assertEqual(list(rows[0]), ["title", "emails", "instagram"])
            self.assertEqual(report["totals"]["emails_added"], 1)

    def test_partial_and_all_failed_reports_are_saved(self):
        for responses, code, status in (([("job1", []), RuntimeError("failure")], 2, "partial"),
                                         ([RuntimeError("failure"), RuntimeError("failure")], 1, "failed")):
            with tempfile.TemporaryDirectory() as root:
                out = str(Path(root, "leads.json"))
                with mock.patch.object(scrape, "req"), mock.patch.object(scrape, "geocode", return_value=("30", "31")), \
                     mock.patch.object(scrape, "run_job", side_effect=responses), redirect_stdout(io.StringIO()), self.assertRaises(SystemExit) as result:
                    scrape.main(["dentists in {city}", "--city", "Cairo", "--city", "Giza", "--no-socials", "--no-email", "--out", out])
                self.assertEqual(result.exception.code, code)
                report = json.loads(Path(root, "leads.report.json").read_text(encoding="utf-8"))
                self.assertEqual(report["status"], status)
                self.assertEqual(len(report["jobs"]), 2)
                self.assertEqual(json.loads(Path(out).read_text(encoding="utf-8")), [])

    def test_geocode_failure_is_reported(self):
        with tempfile.TemporaryDirectory() as root:
            out = str(Path(root, "leads.json"))
            with mock.patch.object(scrape, "req"), mock.patch.object(scrape, "geocode", return_value=None), redirect_stdout(io.StringIO()), self.assertRaises(SystemExit) as result:
                scrape.main(["dentists", "--city", "Unknown", "--out", out])
            self.assertEqual(result.exception.code, 1)
            report = pro.read_json(Path(root, "leads.report.json"))
            self.assertEqual(report["totals"]["locations_failed"], 1)
            self.assertEqual(report["totals"]["jobs_attempted"], 0)

    def test_unknown_active_job_stops_subsequent_jobs(self):
        with tempfile.TemporaryDirectory() as root:
            out = str(Path(root, "leads.json"))
            error = scrape.JobError("polling lost", "active-job-id", active_unknown=True)
            with mock.patch.object(scrape, "req"), mock.patch.object(scrape, "geocode", return_value=("30", "31")), \
                 mock.patch.object(scrape, "run_job", side_effect=error) as jobs, redirect_stdout(io.StringIO()), self.assertRaises(SystemExit):
                scrape.main(["dentists in {city}", "--city", "Cairo", "--city", "Giza", "--no-email", "--no-socials", "--out", out])
            self.assertEqual(jobs.call_count, 1)
            report = pro.read_json(Path(root, "leads.report.json"))
            self.assertEqual(report["totals"]["jobs_not_started"], 1)
            self.assertEqual(report["jobs"][0]["job_id"], "active-job-id")


def minimal_report():
    return {"name": "Egypt campaign", "status": "complete", "started_at": "2026-10-26T10:00:00+00:00", "elapsed_seconds": 3,
            "totals": {"exported_rows": 1}, "collected_coverage": pro.coverage([]), "exported_coverage": pro.coverage([]),
            "locations": [], "jobs": [], "settings": {"languages": ["en", "ar"], "scoring": pro.scoring_rules()}}


class Exports(unittest.TestCase):
    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "optional independent XLSX reader unavailable")
    def test_workbook_with_independent_reader(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as root:
            path = Path(root, "check.xlsx")
            pro.write_results(path, [{"title": "عيادة النور", "phone": "+20100", "lead_score": 90,
                                      "instagram": "https://instagram.com/clinic"}],
                              ["title", "phone", "lead_score", "instagram"], "xlsx", minimal_report(), rtl=True)
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["Summary", "Leads", "Locations", "Jobs", "Settings"])
            sheet = workbook["Leads"]
            self.assertEqual(sheet["A2"].value, "عيادة النور")
            self.assertEqual(sheet["B2"].value, "+20100")
            self.assertEqual(sheet["C2"].value, 90)
            self.assertEqual(sheet["D2"].hyperlink.target, "https://instagram.com/clinic")
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertTrue(sheet.sheet_view.rightToLeft)
            self.assertEqual(sheet.auto_filter.ref, "A1:D2")
            self.assertEqual(len(sheet.conditional_formatting), 1)
            workbook.close()

    def test_xlsx_package_unicode_hyperlinks_styles_and_string_safety(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root, "leads.xlsx")
            fields = ["title", "phone", "review_rating", "lead_score", "instagram"]
            rows = [{"title": "=SUM(1,2)", "phone": "+201001234567", "review_rating": "٤٫٥", "lead_score": 85,
                     "instagram": "https://instagram.com/clinic"}, {"title": "عيادة النور"}]
            pro.write_results(path, rows, fields, "xlsx", minimal_report(), rtl=True)
            with zipfile.ZipFile(path) as archive:
                for name in archive.namelist():
                    ET.fromstring(archive.read(name))
                ns = {"s": pro.NS}
                sheet = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))
                self.assertEqual(sheet.find("s:sheetViews/s:sheetView", ns).get("rightToLeft"), "1")
                self.assertEqual(sheet.find("s:sheetViews/s:sheetView/s:pane", ns).get("state"), "frozen")
                self.assertIsNotNone(sheet.find("s:autoFilter", ns))
                self.assertIsNotNone(sheet.find("s:conditionalFormatting", ns))
                self.assertEqual(sheet.find("s:sheetData/s:row[@r='2']/s:c[@r='A2']", ns).get("t"), "inlineStr")
                self.assertEqual(sheet.find("s:sheetData/s:row[@r='2']/s:c[@r='C2']/s:v", ns).text, "4.5")
                self.assertFalse(sheet.findall(".//s:f", ns))
                self.assertIn("عيادة النور", archive.read("xl/worksheets/sheet2.xml").decode())
                self.assertIn("https://instagram.com/clinic", archive.read("xl/worksheets/_rels/sheet2.xml.rels").decode())

    def test_csv_utf8_bom_and_formula_safety_json_preserves_text(self):
        with tempfile.TemporaryDirectory() as root:
            rows = [{"title": "عيادة", "phone": "+20100", "website": "=HYPERLINK(1)"}]
            fields = list(rows[0])
            csv_path, json_path = Path(root, "nested/leads.csv"), Path(root, "leads.json")
            pro.write_results(csv_path, rows, fields, "csv", minimal_report())
            self.assertTrue(csv_path.read_bytes().startswith(b"\xef\xbb\xbf"))
            with csv_path.open(encoding="utf-8-sig", newline="") as source:
                row = next(csv.DictReader(source))
            self.assertEqual(row["phone"], "'+20100")
            self.assertTrue(row["website"].startswith("'="))
            pro.write_results(json_path, rows, fields, "json", minimal_report())
            self.assertEqual(pro.read_json(json_path), rows)


if __name__ == "__main__":
    unittest.main(verbosity=2)
